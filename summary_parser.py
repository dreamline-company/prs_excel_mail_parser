"""Парсер месячных сводок ПРС (xlsx) в записи для POST /repairs/v1/summaries/parsed.

Устройство книги:
  * лист на каждые сутки — название «20.09.26.»; лист «00.MM.YY.» — шаблон,
    пропускается. Дата записи = дата листа (день, за который сводка).
  * внутри листа четыре секции по НГДУ: Жайыкмунайгаз (без заголовка), далее
    строки-заголовки «Сводка ... по НГДУ "Жылыоймунайгаз" ...» и т.д.
  * блок (одна бригада на скважине) начинается там, где заполнена колонка A —
    это порядковый номер внутри секции, а НЕ номер бригады.

Колонки блока:
  B — мастер и подъёмник с госномером («Утеш А. / ПТП-40 / №237 ALD / 2022 г.в.»);
  C — месторождение, агрегат и бригада («ЮЗК / ДЭЛ-150 / №15585 / Бригада №12»);
  D — номер скважины и способ эксплуатации («251 ШГН», «403R», «1А», «84»);
  N — дата начала ремонта (в запись не идёт);
  O — «скв.№76 20.09.2026» — вторая скважина (откуда переехали);
  P / Q — текст 1-й и 2-й смены.

Имя скважины в бэкенде: ``<КОД_МЕСТОРОЖДЕНИЯ>_<NNNN>`` («BLG_0251»), а с буквенным
суффиксом — ``<КОД>_<NNN><S>`` («ZPV_403R», «BLG_101A»). Суффикс важен: BLG_0101 и
BLG_101A — разные скважины.
"""

from __future__ import annotations

import datetime as dt
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# Справочники
# ---------------------------------------------------------------------------

CAR_PLATE_MAP = {
    "024BH": "024BH06",
    "250AKD": "250AMD",
    "564AZ06": "564AY06",
    "E247AHD": "AHD247E",
    "153AAS": "153AS06",
    "153AC": "153AS06",
    "259ALD06": "ALD259E",
    "480AZD": "AZD480E",
    "730AU": "730AU06",
    "249MD": "249AMD",
}

# Название месторождения (как пишут в сводке, в нижнем регистре) -> код-префикс
# скважины в бэкенде (oil_fields.prefix).
FIELD_CODE_MAP = {
    "акингень": "AKG", "акинген": "AKG",
    "аккудук": "AKD", "аккудык": "AKD",
    "актюбе": "ATB", "актобе": "ATB",
    "алтыколь": "ATK", "алтыкуль": "ATK",
    "б. жоламанов": "JLM", "б жоламанов": "JLM", "б.жоламанов": "JLM",
    "ботахан": "BTN",
    "в. макат": "VMT", "в макат": "VMT", "в.макат": "VMT",
    "вост.макат": "VMT", "вост. макат": "VMT", "вост макат": "VMT", "восточный макат": "VMT",
    "восточный молдабек": "VMB", "ш. молдабек": "VMB", "ш молдабек": "VMB", "ш.молдабек": "VMB",
    "гран": "GRN",
    "досмухамбетовское": "DMB", "досмухамбетиовскоее": "DMB", "досумхамбетовское": "DMB",
    "доссор": "DSR",
    "жанаталап": "ZHT",
    "забурунье": "ZBN", "забурын": "ZBN", "забурынье": "ZBN",
    "западная прорва": "ZPV", "з. прорва": "ZPV", "з прорва": "ZPV", "з.прорва": "ZPV", "зап.прорва": "ZPV", "зап. прорва": "ZPV",
    "карасор западный": "ZKS",
    "каратон": "KRT",
    "карсак": "KRK",
    "кисимбай": "KSB", "кисымбай": "KSB",
    "косшагил": "KSG", "косшагыл": "KSG",
    "кошкар": "KKR",
    "кульсары": "KLR",
    "лиман": "LMN",
    "новобогатинск ю-в": "UVN", "новобогатинск юв": "UVN", "ювн": "UVN",
    "с. балгимбаев": "BLG", "с балгимбаев": "BLG", "с.балгимбаев": "BLG",
    "с. жолдыбай": "SJB", "с жолдыбай": "SJB", "с.жолдыбай": "SJB",
    "с. нуржанов": "NRG", "с нуржанов": "NRG", "с.нуржанов": "NRG", "с.нуржанова": "NRG", "с. нуржанова": "NRG",
    "северный котыртас": "SKS", "с.қотыртас": "SKS", "с қотыртас": "SKS", "с.котыртас": "SKS", "с котыртас": "SKS", "котыртас": "SKS",
    "терен-узек": "TNU", "терен узек": "TNU",
    "уаз": "UAZ", "с. уаз": "UAZ", "с уаз": "UAZ", "с.уаз": "UAZ",
    "уаз восточный": "UZV", "ш.уаз": "UZV",
    "уаз северный": "UZS",
    "ю-в камышитовое": "UVK", "ювк": "UVK",
    "ю-з камышитовое": "UZK", "юзк": "UZK", "южный забурунный купол": "UZK",
}
_FIELD_KEYS_BY_LEN = sorted(FIELD_CODE_MAP, key=len, reverse=True)

# Латинские буквы, которые в сводках печатают вместо похожих кириллических
# («c.уаз» с латинской c) — для сравнения названий месторождений.
_LAT_TO_CYR = str.maketrans("abcehkmoptxy", "авсенкмортху")
# Кириллица -> латиница для госномеров (визуальное соответствие).
_CYR_TO_LAT_PLATE = str.maketrans("АВСЕНКМОРТХУ", "ABCEHKMOPTXY")
# Кириллица -> латиница для суффикса скважины (транслитерация, как в БД:
# «7д» -> VMT_007D, «1А» -> AKG_001A, «6п» -> ..._006P).
_CYR_TO_LAT_SUFFIX = str.maketrans("АБВГДЕЗКНОПРСТУХ", "ABVGDEZKNOPRSTUX")

_PUMP_TYPES = ("УЭЦН", "ЭЦН", "ШГН", "ЭВН", "ЭБС", "ФОНТ", "НАБ")
_PUMP_LABELS = {"ФОНТ": "Фонтан"}

_HAS_TIME = re.compile(r"\d{1,2}[:.]\d{2}\s*[-–]\s*\d{1,2}[:.]\d{2}")
_SHEET_DATE = re.compile(r"^\s*(\d{2})[.\-](\d{2})[.\-](\d{2,4})\.?\s*$")
_WELL = re.compile(r"^\s*(\d{1,4})([A-Za-zА-Яа-я])?(?![A-Za-zА-Яа-я0-9])")
_WELL_SLASH = re.compile(r"^\s*(\d{1,3})\s*/\s*(\d{1,3})(?!\d)")
_BRIGADE_NUM = re.compile(r"(?i)бригада\s*№?\s*(\d{1,3})\b")
_BRIGADE_ANY = re.compile(r"(?i)бригада\s*№?\s*\d*")
_DEVICE_MODEL = re.compile(r"(?i)дэл\s*-?\s*\d+")
_DEVICE_NUM = re.compile(r"№\s*(\d{4,6})\b")
_SECOND_WELL = re.compile(r"(?i)скв\.?\s*№?\s*(\d{1,4})([A-Za-zА-Яа-я])?(?![A-Za-zА-Яа-я0-9])")
_NGDU_HEADER = re.compile(r'(?i)НГДУ\s*[«"]?\s*([^»"\n]+?)\s*[»"]?\s*(?:на\b|$)')
_PLATE = re.compile(r"\b([A-Za-zА-Яа-я]?\s*\d{3}\s*[A-Za-zА-Яа-я]{2,3}\s*(?:/|-)?\s*\d{0,3})\b")
_PHONE = re.compile(r"(?:\+7|8)[\s\-]?\(?[7]\d{2}\)?[\s\-]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}")
_EQUIPMENT = re.compile(r"(?i)(АПРС|ПТП|ПАП|Көтергіш|Подъемник|Сваб|УМК|АЦ|УПСТ)[\s\-]*[\d/]*")
_YEAR_SUFFIX = re.compile(r"\b\d{4}\s*[гж]\.?(\s*в\.?)?")
_TIME_SPLIT = re.compile(r"\s*(?=(?<!\d)\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2})")

EMPTY_MARKERS = ("", "none", "нет", "-", ".", "—")


# ---------------------------------------------------------------------------
# Результат
# ---------------------------------------------------------------------------


@dataclass
class ParseReport:
    file: str
    sheets: int = 0
    skipped_sheets: list[str] = field(default_factory=list)
    blocks: int = 0
    blocks_without_well: int = 0
    blocks_without_shift_text: int = 0
    records: int = 0
    unknown_fields: Counter = field(default_factory=Counter)
    unparsable_wells: Counter = field(default_factory=Counter)
    without_brigade: int = 0
    without_car: int = 0
    per_ngdu: Counter = field(default_factory=Counter)
    inferred_fields: int = 0

    def summary(self) -> str:
        lines = [
            f"{self.file}: листов {self.sheets} (пропущено {len(self.skipped_sheets)}: "
            f"{', '.join(self.skipped_sheets) or '-'}), блоков {self.blocks}, "
            f"без скважины {self.blocks_without_well}, без текста смен "
            f"{self.blocks_without_shift_text}, записей {self.records}",
            f"  по НГДУ: {dict(self.per_ngdu)}",
            f"  без номера бригады: {self.without_brigade}, без госномера: {self.without_car}, "
            f"месторождение взято по машине: {self.inferred_fields}",
        ]
        if self.unknown_fields:
            lines.append(f"  неизвестные месторождения: {dict(self.unknown_fields.most_common(10))}")
        if self.unparsable_wells:
            lines.append(f"  нераспознанные скважины: {dict(self.unparsable_wells.most_common(10))}")
        return "\n".join(lines)


@dataclass
class _Context:
    """Состояние на файл: месторождение по (НГДУ, госномер) для блоков без колонки C."""

    field_by_crew: dict[tuple[str, str], str] = field(default_factory=dict)


@dataclass
class ParseResult:
    records: list[dict]
    report: ParseReport


# ---------------------------------------------------------------------------
# Утилиты
# ---------------------------------------------------------------------------


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\xa0", " ").replace("‑", "-")
    return re.sub(r"[ \t]+", " ", text).strip()


def _is_empty(value: object) -> bool:
    return normalize_text(value).lower() in EMPTY_MARKERS


def sheet_date(title: str) -> dt.date | None:
    """«20.09.26.» -> 2026-09-20; «00.09.26.» (шаблон) и прочее -> None."""
    m = _SHEET_DATE.match(title)
    if not m:
        return None
    day, month, year = (int(x) for x in m.groups())
    if year < 100:
        year += 2000
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def _normalize_field_name(text: str) -> str:
    text = normalize_text(text).lower().replace("ё", "е").replace(":", ".")
    text = re.sub(r"\s*\.\s*", ".", text)
    text = re.sub(r"\.$", "", text)
    return text.strip(" .-")


def field_code(c_cell: object) -> tuple[str | None, str]:
    """Код месторождения из колонки C. Возвращает (код | None, распознанный текст)."""
    text = normalize_text(c_cell)
    text = _BRIGADE_ANY.sub(" ", text)
    text = _DEVICE_MODEL.sub(" ", text)
    text = re.sub(r"№\s*\d+", " ", text)
    lines = [_normalize_field_name(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    if not lines:
        return None, ""
    candidate = lines[0]
    for variant in (candidate, candidate.translate(_LAT_TO_CYR)):
        if variant in FIELD_CODE_MAP:
            return FIELD_CODE_MAP[variant], candidate
        for key in _FIELD_KEYS_BY_LEN:
            if variant.startswith(key) and (len(variant) == len(key) or not variant[len(key)].isalpha()):
                return FIELD_CODE_MAP[key], candidate
    return None, candidate


def well_name(d_cell: object, code: str) -> tuple[str | None, str]:
    """Имя скважины бэкенда из колонки D. Возвращает (имя | None, способ эксплуатации)."""
    text = normalize_text(d_cell).replace("\n", " ")
    upper = text.upper()
    pump = "Не указан"
    for token in _PUMP_TYPES:
        if token in upper:
            pump = _PUMP_LABELS.get(token, token)
            break
    m_slash = _WELL_SLASH.match(text)
    if m_slash:
        return f"{code}_{m_slash.group(1)}/{m_slash.group(2)}", pump
    m = _WELL.match(text)
    if not m:
        return None, pump
    number = int(m.group(1))
    suffix = (m.group(2) or "").upper().translate(_CYR_TO_LAT_SUFFIX)
    if suffix:
        name = f"{code}_{number:03d}{suffix}" if number < 1000 else f"{code}_{number}{suffix}"
    else:
        name = f"{code}_{number:04d}"
    return name, pump


def brigade_number(c_cell: object, b_cell: object = None) -> int:
    for cell in (c_cell, b_cell):
        m = _BRIGADE_NUM.search(normalize_text(cell))
        if m:
            return int(m.group(1))
    return 0


def device_number(c_cell: object) -> str:
    text = _BRIGADE_ANY.sub(" ", normalize_text(c_cell))
    m = _DEVICE_NUM.search(text)
    return m.group(1) if m else ""


def second_well_name(o_cell: object, code: str) -> str:
    text = normalize_text(o_cell).replace("\n", " ")
    m = _SECOND_WELL.search(text)
    if not m:
        return ""
    number = int(m.group(1))
    if 2000 <= number <= 2099 and not m.group(2):
        return ""  # это год из даты, а не скважина
    name, _ = well_name(f"{m.group(1)}{m.group(2) or ''}", code)
    return name or ""


def normalize_plate(plate: str) -> str:
    return plate.upper().translate(_CYR_TO_LAT_PLATE)


def car_and_master(b_cell: object) -> tuple[str, str]:
    """«Фамилия И. 249AMD» из колонки B. Возвращает (значение car, госномер)."""
    raw = normalize_text(b_cell)
    search_str = re.sub(r"(?i)kz\s*", " ", raw)

    candidates: list[tuple[str, bool, str]] = []
    for m in _PLATE.finditer(search_str):
        preceding = search_str[max(0, m.start() - 35) : m.start()].upper()
        priority = any(marker in preceding for marker in ("АПРС", "ПТП", "ПАП"))
        clean = re.sub(r"[\s/|\-]", "", m.group(1))
        candidates.append((clean, priority, m.group(1)))

    plate = ""
    if candidates:
        candidates.sort(key=lambda x: x[1], reverse=True)
        plate = candidates[0][0]
    else:
        for line in search_str.splitlines():
            if "№" in line:
                plate = line.replace("№", "").replace(" ", "").strip()
                candidates.append((plate, False, line))
                break
    if plate:
        plate = normalize_plate(plate)
        plate = CAR_PLATE_MAP.get(plate, plate)

    lines = [x.strip() for x in raw.splitlines() if x.strip()]
    name_line = lines[0] if lines else ""
    for _, _, raw_plate in candidates:
        name_line = name_line.replace(raw_plate, " ")
    name_line = _PHONE.sub(" ", name_line)
    name_line = _EQUIPMENT.sub(" ", name_line)
    name_line = _YEAR_SUFFIX.sub(" ", name_line)
    name_line = re.sub(r"(?i)kz\s*", " ", name_line)
    name_line = _BRIGADE_ANY.sub(" ", name_line).replace("№", "")
    name_line = re.sub(r"(?i)\b\d\s*-?\s*смена\b", " ", name_line)
    name_line = re.sub(r"(?i)^(день|ночь)\s*[-–]?\s*", "", normalize_text(name_line))
    parts = [
        p for p in name_line.split()
        if any(ch.isalpha() for ch in p) and p.lower() not in ("мастер", "ст.мастер", "ст.", "ст", "смена", "прс")
    ]
    master = ""
    if len(parts) >= 2:
        master = f"{parts[0].capitalize()} {parts[1][0].upper()}."
    elif len(parts) == 1:
        master = parts[0].capitalize()

    if master and plate:
        return f"{master} {plate}", plate
    if plate:
        return plate, plate
    if master:
        return master, ""
    return "Не указана", ""


def parse_shift_text(raw_text: object) -> list[str]:
    text = normalize_text(raw_text)
    if text.lower() in EMPTY_MARKERS or not _HAS_TIME.search(text):
        return []
    text = text.replace("::", ":")
    parts: list[str] = []
    for line in text.split("\n"):
        parts.extend(_TIME_SPLIT.split(line))
    result: list[str] = []
    for part in parts:
        cleaned = normalize_text(part)
        cleaned = re.sub(r"^[Вв]р\s*\.?\s*", "", cleaned)
        cleaned = re.sub(r"\s*[Вв]р\s*\.?\s*$", "", cleaned)
        cleaned = normalize_text(cleaned).replace('"', "'")
        if cleaned:
            if not cleaned.endswith("."):
                cleaned += "."
            result.append(cleaned)
    return result


def _ngdu_from_header(a_cell: object) -> str | None:
    text = normalize_text(a_cell)
    if "нгду" not in text.lower():
        return None
    m = _NGDU_HEADER.search(text)
    return m.group(1).strip() if m else "?"


# ---------------------------------------------------------------------------
# Разбор книги
# ---------------------------------------------------------------------------


def _as_int(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = normalize_text(value)
    return int(text) if text.isdigit() else None


def parse_sheet(ws, day: dt.date, report: ParseReport, ctx: _Context) -> list[dict]:  # noqa: ANN001
    records: list[dict] = []
    ngdu = "Жайыкмунайгаз"
    block: dict | None = None

    def flush() -> None:
        nonlocal block
        if block is None:
            return
        _emit_block(block, day, ngdu_name=block["ngdu"], report=report, out=records, ctx=ctx)
        block = None

    for row in range(4, ws.max_row + 1):
        a = ws.cell(row=row, column=1).value
        header_ngdu = _ngdu_from_header(a)
        if header_ngdu is not None:
            flush()
            ngdu = header_ngdu
            continue
        ordinal = _as_int(a)
        cells = {
            col: ws.cell(row=row, column=idx).value
            for col, idx in (("B", 2), ("C", 3), ("D", 4), ("N", 14), ("O", 15), ("P", 16), ("Q", 17))
        }
        if ordinal is not None:
            flush()
            block = {"ngdu": ngdu, **cells}
            continue
        if isinstance(a, str) and a.strip():
            # служебная строка («Всего потери по дебитам ...»)
            flush()
            continue
        if block is None:
            continue
        # продолжение блока (объединённые ячейки): дописываем текст смен
        for col in ("P", "Q"):
            extra = normalize_text(cells[col])
            if extra and extra.lower() not in EMPTY_MARKERS:
                block[col] = f"{normalize_text(block[col])}\n{extra}".strip()
    flush()
    return records


def _emit_block(block: dict, day: dt.date, *, ngdu_name: str, report: ParseReport, out: list[dict], ctx: _Context) -> None:
    report.blocks += 1
    if _is_empty(block["D"]):
        report.blocks_without_well += 1
        return
    shift_1 = parse_shift_text(block["P"])
    shift_2 = parse_shift_text(block["Q"])
    if not shift_1 and not shift_2:
        report.blocks_without_shift_text += 1
        return

    car, plate = car_and_master(block["B"])
    code, field_text = field_code(block["C"])
    crew_key = (ngdu_name, plate)
    if code is None and not field_text and plate and crew_key in ctx.field_by_crew:
        code = ctx.field_by_crew[crew_key]
        report.inferred_fields += 1
    if code is None:
        report.unknown_fields[f"{ngdu_name}: {field_text or '<пусто>'}"] += 1
        return
    if plate and field_text:
        ctx.field_by_crew[crew_key] = code
    name, pump = well_name(block["D"], code)
    if name is None:
        report.unparsable_wells[f"{code}: {normalize_text(block['D'])[:20]}"] += 1
        return

    brigade = brigade_number(block["C"], block["B"])
    if not brigade:
        report.without_brigade += 1
    if not plate:
        report.without_car += 1

    base = {
        "start_date": day.strftime("%d.%m.%Y"),
        "brigade_number": brigade,
        "well_name": name,
        "second_well_name": second_well_name(block["O"], code),
        "pump_type": pump,
        "car": car,
        "device_number": device_number(block["C"]),
    }
    for shift_no, details in ((1, shift_1), (2, shift_2)):
        if details:
            out.append({**base, "shift_type_number": shift_no, "shift_details": details})
            report.records += 1
            report.per_ngdu[ngdu_name] += 1


def parse_workbook(path: str | Path) -> ParseResult:
    path = Path(path)
    report = ParseReport(file=path.name)
    records: list[dict] = []
    ctx = _Context()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        # листы идут от конца месяца к началу — разбираем по возрастанию даты,
        # чтобы месторождение бригады «наследовалось» вперёд по времени
        sheets = sorted(wb.worksheets, key=lambda w: sheet_date(w.title) or dt.date.min)
        for ws in sheets:
            day = sheet_date(ws.title)
            if day is None:
                report.skipped_sheets.append(ws.title)
                continue
            report.sheets += 1
            records.extend(parse_sheet(ws, day, report, ctx))
    finally:
        wb.close()
    return ParseResult(records=records, report=report)
